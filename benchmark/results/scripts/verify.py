#!/usr/bin/env python3
"""Verify the public wide CSV, dictionary, manifest, and XLSX cell by cell."""
import argparse
import csv
import datetime as dt
import hashlib
import itertools
import json
import math
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook


def sha(path):
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1 << 20), b""): h.update(chunk)
    return h.hexdigest()


def convert(value, kind):
    if value == "": return None
    if kind == "bool": return value.lower() == "true"
    if kind == "int": return int(value)
    if kind == "float": return float(value)
    if kind == "timestamp": return dt.datetime.fromisoformat(value.replace("Z", "+00:00")).astimezone(dt.timezone.utc).replace(tzinfo=None)
    return "'" + value if value.startswith("=") else value


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--result", type=Path, required=True); parser.add_argument("--report", type=Path)
    args = parser.parse_args(); root = args.result.resolve(); errors = []; checks = []
    manifest = json.loads((root / "manifest.json").read_text(encoding="utf-8"))
    for item in manifest["outputs"]:
        if not (root / item["file"]).is_file() or sha(root / item["file"]) != item["sha256"]: errors.append("hash:" + item["file"])
    checks.append({"name": "output_hashes", "passed": not errors, "count": len(manifest["outputs"])})
    with (root / "dictionary.csv").open(encoding="utf-8", newline="") as stream: dictionary = list(csv.DictReader(stream))
    fields = [row["field"] for row in dictionary]; meta = {row["field"]: row for row in dictionary}
    with (root / "measurements.csv").open(encoding="utf-8", newline="") as stream:
        reader = csv.DictReader(stream); rows = list(reader)
    if reader.fieldnames != fields: errors.append("dictionary/order")
    if len(rows) != manifest["counts"]["rows"]: errors.append("row/count")
    ids = [row["record_id"] for row in rows]
    if len(ids) != len(set(ids)): errors.append("record_id/duplicate")
    cells = {row["cell_id"] for row in rows if row["cell_id"]}
    anchors = [row for row in rows if row["is_cell_anchor"] == "true"]
    if len(anchors) != manifest["counts"]["cell_anchors"] or len({row["cell_id"] for row in anchors}) != len(anchors): errors.append("anchor/contract")
    warmups = [row for row in rows if row["is_warmup"] == "true"]
    if any(row["round"] != "0" or row["include_in_default_stats"] != "false" for row in warmups): errors.append("warmup/contract")
    workbook = load_workbook(root / "measurements.xlsx", read_only=True, data_only=False)
    if workbook.sheetnames != ["Описание", "Измерения", "Словарь"]: errors.append("xlsx/sheets")
    ws = workbook["Измерения"]; iterator = ws.iter_rows(); header = [cell.value for cell in next(iterator)]
    if header != fields: errors.append("xlsx/header")
    value_count = 0; sheet_row_count = 0; sentinel = object()
    for row_no, pair in enumerate(itertools.zip_longest(iterator, rows, fillvalue=sentinel), 2):
        sheet_row, csv_row = pair
        if sheet_row is sentinel or csv_row is sentinel:
            errors.append("xlsx/row_count"); break
        sheet_row_count += 1
        for col_no, (cell, name) in enumerate(zip(sheet_row, fields), 1):
            expected = convert(csv_row[name], meta[name]["type"]); actual = cell.value; value_count += 1
            if isinstance(expected, float): ok = isinstance(actual, (int, float)) and math.isclose(expected, actual, rel_tol=1e-12, abs_tol=1e-12)
            else: ok = actual == expected
            if not ok and len(errors) < 50: errors.append(f"xlsx/value:{row_no}:{col_no}:{name}")
    if sheet_row_count != len(rows): errors.append("xlsx/row_count")
    workbook.close()
    with zipfile.ZipFile(root / "measurements.xlsx") as archive:
        ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        formulas = errors_in_cells = 0
        for name in archive.namelist():
            if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
                xml = ET.fromstring(archive.read(name)); formulas += len(xml.findall(".//m:f", ns)); errors_in_cells += sum(c.get("t") == "e" for c in xml.findall(".//m:c", ns))
        if formulas: errors.append("xlsx/formulas")
        if errors_in_cells: errors.append("xlsx/errors")
    checks.extend([
        {"name": "row_keys_and_grain", "passed": not any(x.startswith(("record_id", "anchor", "warmup", "row/")) for x in errors), "rows": len(rows), "cells": len(cells)},
        {"name": "xlsx_values", "passed": not any(x.startswith("xlsx/") for x in errors), "values_checked": value_count},
    ])
    report = {"summary": {"passed": sum(x["passed"] for x in checks), "failed": sum(not x["passed"] for x in checks)}, "checks": checks, "errors": errors, "xlsx_sha256": sha(root / "measurements.xlsx")}
    target = args.report or root / "acceptance.json"; target.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report["summary"] | {"rows": len(rows), "xlsx_values": value_count}, ensure_ascii=False))
    if errors: raise SystemExit(1)


if __name__ == "__main__": main()
